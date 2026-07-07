"""
excel_service.py
----------------
엑셀 파일(.xlsx)을 데이터베이스처럼 다루는 저수준 서비스.

- 엑셀 파일 열기 / 저장 (없으면 생성)
- 시트 확인 및 자동 생성
- 헤더 확인 및 자동 생성
- 시트 → DataFrame 읽기
- DataFrame → 시트 쓰기
- 자유 형식 그리드(list of list) → 시트 쓰기 (보고서용)

openpyxl Workbook 객체를 단일 원본으로 사용하여 기존 시트를 최대한 보존한다.
"""

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from . import backup_service, config


# ---------------------------------------------------------------------------
# 워크북 열기 / 저장
# ---------------------------------------------------------------------------

def load_or_create_workbook(excel_path: Path = None) -> Workbook:
    """
    엑셀 파일을 연다. 없으면 새 워크북을 만들고 기본 구조를 채운다.
    항상 모든 필수 시트/헤더가 준비된 상태의 Workbook 을 반환한다.
    """
    excel_path = Path(excel_path) if excel_path else config.EXCEL_PATH

    if excel_path.exists():
        wb = load_workbook(excel_path)
    else:
        wb = Workbook()
        # 기본 생성되는 시트 제거 (아래에서 필요한 시트를 새로 만든다)
        default = wb.active
        wb.remove(default)

    ensure_structure(wb)
    return wb


def new_workbook() -> Workbook:
    """디스크를 읽지 않고 필수 구조만 갖춘 새 워크북을 만든다(테스트/시드용)."""
    wb = Workbook()
    wb.remove(wb.active)
    ensure_structure(wb)
    return wb


def save_workbook(wb: Workbook, excel_path: Path = None) -> Path:
    """워크북을 지정 경로에 저장한다. 상위 폴더가 없으면 생성한다."""
    excel_path = Path(excel_path) if excel_path else config.EXCEL_PATH
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return excel_path


def save_with_backup(wb: Workbook, excel_path: Path = None):
    """
    파일을 저장하기 전에 현재 디스크 상의 파일을 백업한 뒤 저장한다.
    거래 저장/수정/삭제, 보고서 갱신 등 파일을 변경하는 모든 작업에서 사용한다.
    반환: (저장 경로, 백업 경로 또는 None)
    """
    excel_path = Path(excel_path) if excel_path else config.EXCEL_PATH
    backup_path = backup_service.create_backup(excel_path)
    saved = save_workbook(wb, excel_path)
    return saved, backup_path


# ---------------------------------------------------------------------------
# 구조(시트/헤더) 보장
# ---------------------------------------------------------------------------

def ensure_structure(wb: Workbook) -> Workbook:
    """
    필수 시트가 없으면 생성하고, 시트 순서를 정렬한다.
    거래입력대장에는 헤더가 없으면 헤더 행을 만든다.
    """
    for name in config.ALL_SHEETS:
        if name not in wb.sheetnames:
            wb.create_sheet(title=name)

    # 거래입력대장 헤더 보장
    ensure_header(wb, config.SHEET_LEDGER, config.LEDGER_COLUMNS)

    # 시트 순서를 config.ALL_SHEETS 순서로 정렬 (그 외 시트는 뒤에 유지)
    desired = [n for n in config.ALL_SHEETS if n in wb.sheetnames]
    others = [n for n in wb.sheetnames if n not in config.ALL_SHEETS]
    wb._sheets.sort(key=lambda ws: (desired + others).index(ws.title))
    return wb


def ensure_header(wb: Workbook, sheet_name: str, columns: list[str]) -> None:
    """
    시트 첫 행이 비어있으면 columns 를 헤더로 기록한다.
    이미 값이 있으면 건드리지 않는다(기존 시트 보존).
    """
    if sheet_name not in wb.sheetnames:
        wb.create_sheet(title=sheet_name)
    ws = wb[sheet_name]

    first_cell = ws.cell(row=1, column=1).value
    if first_cell is None or str(first_cell).strip() == "":
        for idx, col in enumerate(columns, start=1):
            ws.cell(row=1, column=idx, value=col)


# ---------------------------------------------------------------------------
# 셀 값 정규화
# ---------------------------------------------------------------------------

def _normalize_cell(value):
    """엑셀 셀 값을 파이썬/문자열 친화적으로 정규화한다."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        # 시간이 00:00:00 이면 날짜만
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return value


# ---------------------------------------------------------------------------
# 시트 → DataFrame
# ---------------------------------------------------------------------------

def read_sheet_df(wb: Workbook, sheet_name: str) -> pd.DataFrame:
    """
    시트를 DataFrame 으로 읽는다. 첫 행을 헤더로 사용한다.
    빈 시트이면 빈 DataFrame 을 반환한다.
    """
    if sheet_name not in wb.sheetnames:
        return pd.DataFrame()

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()

    header = [(_normalize_cell(c) if c is not None else f"col{i}")
              for i, c in enumerate(rows[0])]
    data = []
    for r in rows[1:]:
        # 완전히 빈 행은 건너뛴다
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        data.append([_normalize_cell(c) for c in r])

    df = pd.DataFrame(data, columns=header)
    return df


def read_ledger_df(wb: Workbook) -> pd.DataFrame:
    """
    거래입력대장을 표준 컬럼/타입으로 읽는다.
    - 누락 컬럼은 빈 값으로 채운다
    - 금액은 숫자형으로 변환
    - 회계연도는 정수 문자열/숫자를 정수로 변환
    """
    df = read_sheet_df(wb, config.SHEET_LEDGER)

    # 컬럼 보정
    if df.empty:
        df = pd.DataFrame(columns=config.LEDGER_COLUMNS)
    for col in config.LEDGER_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[config.LEDGER_COLUMNS]  # 순서 고정

    # 타입 변환
    df["금액"] = pd.to_numeric(df["금액"], errors="coerce").fillna(0)
    df["회계연도"] = pd.to_numeric(df["회계연도"], errors="coerce").fillna(0).astype(int)

    # 문자열 컬럼 정리
    for col in df.columns:
        if col not in ("금액", "회계연도"):
            df[col] = df[col].astype(str).replace("nan", "").fillna("")

    return df


# ---------------------------------------------------------------------------
# DataFrame / 그리드 → 시트
# ---------------------------------------------------------------------------

def _clear_sheet(ws) -> None:
    """시트의 모든 값을 지운다(행 삭제)."""
    if ws.max_row >= 1:
        ws.delete_rows(1, ws.max_row)


def write_df_to_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    """
    DataFrame 을 시트에 기록한다(기존 내용 삭제 후 헤더 + 데이터 작성).
    """
    if sheet_name not in wb.sheetnames:
        wb.create_sheet(title=sheet_name)
    ws = wb[sheet_name]
    _clear_sheet(ws)

    # 헤더
    for c_idx, col in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c_idx, value=str(col))

    # 데이터
    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for c_idx, col in enumerate(df.columns, start=1):
            ws.cell(row=r_idx, column=c_idx, value=_to_cell(row[col]))

    _auto_width(ws, df.columns)


def write_grid_to_sheet(wb: Workbook, sheet_name: str, grid: list[list]) -> None:
    """
    자유 형식 2차원 그리드를 시트에 기록한다(보고서 레이아웃용).
    """
    if sheet_name not in wb.sheetnames:
        wb.create_sheet(title=sheet_name)
    ws = wb[sheet_name]
    _clear_sheet(ws)

    max_cols = 1
    for r_idx, row in enumerate(grid, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=_to_cell(val))
        max_cols = max(max_cols, len(row))

    # 열 너비 대략 설정
    for c_idx in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 22


def _to_cell(value):
    """DataFrame/파이썬 값을 엑셀에 안전하게 쓸 수 있는 형태로 변환한다."""
    if value is None:
        return ""
    # pandas NaN 처리
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    # numpy 정수/실수 → 파이썬 기본형
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            return value
    return value


def _auto_width(ws, columns) -> None:
    """열 너비를 헤더 길이에 맞춰 대략 설정한다."""
    for c_idx, col in enumerate(columns, start=1):
        width = max(10, min(40, len(str(col)) * 2 + 4))
        ws.column_dimensions[get_column_letter(c_idx)].width = width
