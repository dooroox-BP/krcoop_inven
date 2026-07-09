# -*- coding: utf-8 -*-
"""
================================================================================
 렌탈/케어 서비스 정산 자동화 프로그램  (Config/Data 분리 구조)
================================================================================
'정산자동화_통합서식_V2.xlsx' (아래 2개 시트) 를 읽어 정산을 자동 계산한다.
    - '단가표_설정' 시트 : 작업분류_항목 → 계약단가(원)  (설정/Config)
    - '정산데이터'  시트 : 실제 방문/작업 내역            (데이터/Data)

계산 항목
    (A) 수수료        : 작업분류(선택) → 단가표 매핑
    (B) 추가(통행료)  : ① 주소에 '영종' 포함  → +11,000원
                        ② 카카오 길찾기 왕복거리 50km 이상 → +20,000원 (중복 합산)
    (C) 현장결제      : 기존 입력 유지, 빈칸은 0

결과
    '2026_05_정산완료(자동화).xlsx'
      · 각 행에 (A)(B)(C) 및 왕복거리(km) 기록
      · [합계행]   A 총합 / B 총합 / C 총합
      · [청구액행] "총 청구액 = (A + B) - C" 와 최종 금액(예: 835,000원)
                   → 최종 금액 셀은 노란색 배경

--------------------------------------------------------------------------------
[설치]  터미널에서 아래 한 줄 실행
    pip install pandas openpyxl requests

[사용법]
    1) 아래 KAKAO_API_KEY 에 카카오 REST API 키를 입력한다.
       (키가 없으면 거리 할증은 자동으로 건너뛰고, 영종 할증/수수료만 계산)
    2) '정산자동화_통합서식_V2.xlsx' 를 이 스크립트와 같은 폴더에 둔다.
    3) 실행:
           python rental_settlement.py

    (선택) 테스트용 샘플 입력 파일을 만들고 싶다면:
           python rental_settlement.py --make-sample
           → '정산자동화_통합서식_V2.xlsx' 샘플 생성 후, 위 3) 명령을 그대로 실행

    (선택) 입력/출력 경로 지정:
           python rental_settlement.py --input 원본.xlsx --output 결과.xlsx
================================================================================
"""

import os
import sys
import time
import numbers
import argparse

import pandas as pd
import requests
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =====================================================================
# 0. 기본 설정 (Config)
# =====================================================================
# ▼▼▼ 카카오 REST API 키를 여기에 입력하세요 ▼▼▼
KAKAO_API_KEY = "YOUR_API_KEY_HERE"
# ▲▲▲ (미설정 시 거리 할증은 생략됩니다) ▲▲▲

INPUT_FILE  = "정산자동화_통합서식_V2.xlsx"
OUTPUT_FILE = "2026_05_정산완료(자동화).xlsx"

SHEET_CONFIG = "단가표_설정"   # 설정(단가) 시트
SHEET_DATA   = "정산데이터"    # 데이터 시트
SHEET_OUT    = "정산완료"      # 출력 시트명

# 고정 출발지 (꿈담빌딩)
ORIGIN_ADDRESS = "인천광역시 미추홀구 아암대로 45번길 77-8"

# --- 컬럼명 (실제 엑셀 헤더가 다르면 이 부분만 수정) ---
# [단가표_설정] 시트
COL_PRICE_ITEM  = "작업분류_항목"   # 딕셔너리 Key
COL_PRICE_VALUE = "계약단가(원)"    # 딕셔너리 Value
# [정산데이터] 시트
COL_CATEGORY = "작업분류(선택)"     # H열: 수수료 매핑 기준
COL_FEE      = "수수료"             # (A)
COL_ADDRESS  = "주소"
COL_TOLL     = "추가(통행료)"       # (B)
COL_ROUND    = "왕복거리(km)"
COL_ONSITE   = "현장결제"           # (C)
COL_FILTERS  = ["필터1", "필터2", "필터3"]   # 작업분류 미매칭 시 보조 스캔 대상

# --- 할증 규칙 ---
YEONGJONG_KEYWORD          = "영종"
YEONGJONG_SURCHARGE        = 11_000
LONG_DISTANCE_THRESHOLD_KM = 50.0
LONG_DISTANCE_SURCHARGE    = 20_000

# --- 카카오 API 엔드포인트 ---
KAKAO_ADDR_URL    = "https://dapi.kakao.com/v2/local/search/address.json"    # 주소 → 좌표
KAKAO_KEYWORD_URL = "https://dapi.kakao.com/v2/local/search/keyword.json"    # 키워드 → 좌표(폴백)
KAKAO_NAVI_URL    = "https://apis-navi.kakaomobility.com/v1/directions"      # 자동차 길찾기

API_SLEEP = 0.2   # 연속 호출 사이 지연(초) - 쿼터 보호용


# =====================================================================
# 1. 카카오 API 클라이언트 (좌표 변환 / 왕복거리)
# =====================================================================
_session        = requests.Session()
_geocode_cache  = {}   # 주소 → (x, y)   (동일 주소 재호출 방지)
_distance_cache = {}   # (출발, 도착) → 편도 km


def _api_ready() -> bool:
    """유효한 API 키가 설정되어 있는지 확인."""
    return bool(KAKAO_API_KEY) and KAKAO_API_KEY != "YOUR_API_KEY_HERE"


def _kakao_headers() -> dict:
    return {"Authorization": f"KakaoAK {KAKAO_API_KEY}"}


def geocode(address):
    """
    주소 텍스트 → (경도 x, 위도 y) 튜플. 실패하면 None 을 돌려준다.
    1차로 주소 검색, 결과가 없으면 키워드 검색(건물명/상호)으로 폴백한다.
    * 예외처리: 좌표를 못 찾아도 에러를 던지지 않고 None 반환.
    """
    if address is None or (isinstance(address, float) and pd.isna(address)):
        return None
    key = str(address).strip()
    if not key:
        return None
    if key in _geocode_cache:
        return _geocode_cache[key]

    result = None
    if _api_ready():
        try:
            # 1) 주소 검색
            r = _session.get(KAKAO_ADDR_URL, headers=_kakao_headers(),
                             params={"query": key}, timeout=10)
            docs = r.json().get("documents", []) if r.ok else []
            # 2) 키워드 검색 폴백
            if not docs:
                time.sleep(API_SLEEP)
                r = _session.get(KAKAO_KEYWORD_URL, headers=_kakao_headers(),
                                 params={"query": key}, timeout=10)
                docs = r.json().get("documents", []) if r.ok else []
            if docs:
                result = (float(docs[0]["x"]), float(docs[0]["y"]))
            time.sleep(API_SLEEP)
        except Exception as e:                      # 네트워크/파싱 오류 등 전부 흡수
            print(f"    [경고] 좌표 변환 실패 → 건너뜀: '{key}' ({e})")
            result = None

    _geocode_cache[key] = result
    return result


def driving_distance_km(origin_xy, dest_xy):
    """
    출발지→도착지 자동차 추천경로 편도 주행거리(km). 실패하면 None.
    카카오모빌리티 길찾기 API 사용.
    """
    if origin_xy is None or dest_xy is None:
        return None
    cache_key = (origin_xy, dest_xy)
    if cache_key in _distance_cache:
        return _distance_cache[cache_key]

    dist_km = None
    if _api_ready():
        try:
            params = {
                "origin":      f"{origin_xy[0]},{origin_xy[1]}",
                "destination": f"{dest_xy[0]},{dest_xy[1]}",
                "priority":    "RECOMMEND",   # 추천 경로
            }
            r = _session.get(KAKAO_NAVI_URL, headers=_kakao_headers(),
                             params=params, timeout=10)
            if r.ok:
                routes = r.json().get("routes", [])
                if routes and routes[0].get("result_code", -1) == 0:
                    meters = routes[0]["summary"]["distance"]
                    dist_km = meters / 1000.0
            time.sleep(API_SLEEP)
        except Exception as e:
            print(f"    [경고] 경로 거리 계산 실패 → 건너뜀 ({e})")
            dist_km = None

    _distance_cache[cache_key] = dist_km
    return dist_km


# =====================================================================
# 2. 유틸리티
# =====================================================================
def to_number(value):
    """빈칸/NaN → 0, '12,000원' 같은 문자열 → 12000 으로 정규화."""
    try:
        if pd.isna(value):
            return 0
    except (TypeError, ValueError):
        pass
    if isinstance(value, numbers.Number):
        f = float(value)
        return int(f) if f.is_integer() else f
    s = str(value).strip().replace(",", "").replace("원", "").replace(" ", "")
    if s == "":
        return 0
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return 0


# =====================================================================
# 3. [Step 1] 단가표_설정 → 딕셔너리
# =====================================================================
def load_price_map(path: str) -> dict:
    """'단가표_설정' 시트를 읽어 {작업분류_항목: 계약단가(원)} 딕셔너리로 변환."""
    df = pd.read_excel(path, sheet_name=SHEET_CONFIG)
    if COL_PRICE_ITEM not in df.columns or COL_PRICE_VALUE not in df.columns:
        raise KeyError(
            f"'{SHEET_CONFIG}' 시트에서 '{COL_PRICE_ITEM}' 또는 "
            f"'{COL_PRICE_VALUE}' 컬럼을 찾을 수 없습니다. 실제 컬럼: {list(df.columns)}"
        )
    price_map = {}
    for _, row in df.iterrows():
        item = row[COL_PRICE_ITEM]
        if pd.isna(item):
            continue
        price_map[str(item).strip()] = to_number(row[COL_PRICE_VALUE])
    return price_map


# =====================================================================
# 4. [Step 2] 수수료(A) 조회
# =====================================================================
def lookup_fee(row, price_map: dict):
    """
    작업분류(선택) 값을 단가표에서 찾아 수수료를 반환.
    미매칭 시 보조 로직으로 필터1~3 텍스트를 스캔, 그래도 없으면 0.
    """
    # 1) 기본: 작업분류(선택) 기준
    category = row.get(COL_CATEGORY)
    if category is not None and not pd.isna(category):
        key = str(category).strip()
        if key and key in price_map:
            return price_map[key]

    # 2) 보조: 필터1~3 스캔
    for fcol in COL_FILTERS:
        if fcol in row.index and not pd.isna(row[fcol]):
            fkey = str(row[fcol]).strip()
            if fkey and fkey in price_map:
                return price_map[fkey]

    # 3) 미매칭
    return 0


# =====================================================================
# 5. [Step 2~3] 정산데이터 계산
# =====================================================================
def process(path: str, price_map: dict) -> pd.DataFrame:
    global COL_CATEGORY
    df = pd.read_excel(path, sheet_name=SHEET_DATA)

    # 작업분류(선택) 컬럼명이 다르면 H열(8번째) 위치로 폴백 인식
    if COL_CATEGORY not in df.columns and len(df.columns) >= 8:
        COL_CATEGORY = df.columns[7]
        print(f"    [안내] '{COL_CATEGORY}' 를 작업분류(H열) 컬럼으로 인식")

    # 고정 출발지 좌표 1회 조회
    origin_xy = geocode(ORIGIN_ADDRESS)
    if origin_xy is None and _api_ready():
        print("    [안내] 출발지 좌표 조회 실패 → 거리 할증 생략")

    fees, tolls, rounds = [], [], []
    for _, row in df.iterrows():
        # ---------- (A) 수수료 ----------
        fee = lookup_fee(row, price_map)

        # ---------- (B) 추가(통행료) ----------
        toll = 0
        addr = row.get(COL_ADDRESS)
        addr_str = "" if (addr is None or pd.isna(addr)) else str(addr)

        # ① 지역 할증: 주소에 '영종' 포함
        if YEONGJONG_KEYWORD in addr_str:
            toll += YEONGJONG_SURCHARGE

        # ② 거리 할증: 왕복 50km 이상
        round_km = ""   # 좌표를 못 구하면 빈칸 유지
        if origin_xy is not None and addr_str.strip():
            dest_xy = geocode(addr_str)
            one_way = driving_distance_km(origin_xy, dest_xy)
            if one_way is not None:
                round_km = round(one_way * 2, 1)    # 왕복 = 편도 × 2 (소수점 1자리)
                if round_km >= LONG_DISTANCE_THRESHOLD_KM:
                    toll += LONG_DISTANCE_SURCHARGE   # 영종 할증과 중복 합산

        fees.append(fee)
        tolls.append(toll)
        rounds.append(round_km)

    # 계산 결과를 각 컬럼에 기록 (기존 컬럼이 있으면 값만 덮어씀 = 위치 유지)
    df[COL_FEE]   = fees
    df[COL_TOLL]  = tolls
    df[COL_ROUND] = rounds

    # ---------- (C) 현장결제: 기존값 유지, 빈칸 0 ----------
    if COL_ONSITE in df.columns:
        df[COL_ONSITE] = df[COL_ONSITE].apply(to_number)
    else:
        df[COL_ONSITE] = 0

    return df


# =====================================================================
# 6. [Step 4] 합계행 + 총 청구액행 작성 & 저장 (openpyxl 서식)
# =====================================================================
def save_with_summary(df: pd.DataFrame, output_file: str):
    # A / B / C 총합
    sum_a = int(pd.to_numeric(df[COL_FEE],    errors="coerce").fillna(0).sum())
    sum_b = int(pd.to_numeric(df[COL_TOLL],   errors="coerce").fillna(0).sum())
    sum_c = int(pd.to_numeric(df[COL_ONSITE], errors="coerce").fillna(0).sum())
    final_amount = (sum_a + sum_b) - sum_c   # 총 청구액 = (A + B) - C

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=SHEET_OUT, index=False)
        ws = writer.sheets[SHEET_OUT]

        cols = list(df.columns)
        col_idx = lambda name: cols.index(name) + 1   # 1-based 컬럼 인덱스

        a_col = col_idx(COL_FEE)
        b_col = col_idx(COL_TOLL)
        c_col = col_idx(COL_ONSITE)
        r_col = col_idx(COL_ROUND)
        n_cols = len(cols)

        n_rows     = len(df)
        first_data = 2
        last_data  = 1 + n_rows        # 헤더(1행) 다음부터 데이터
        total_row  = last_data + 1     # 합계행
        bill_row   = total_row + 1     # 총 청구액행

        money_fmt = "#,##0"
        bold      = Font(bold=True)
        thin_top  = Border(top=Side(style="thin"))

        # --- 데이터 영역 숫자 서식 ---
        for cc in (a_col, b_col, c_col):
            for rr in range(first_data, last_data + 1):
                ws.cell(row=rr, column=cc).number_format = money_fmt
        for rr in range(first_data, last_data + 1):
            ws.cell(row=rr, column=r_col).number_format = "0.0"

        # --- 헤더 강조 ---
        header_fill = PatternFill("solid", fgColor="D9E1F2")
        for cc in range(1, n_cols + 1):
            hc = ws.cell(row=1, column=cc)
            hc.font = bold
            hc.fill = header_fill
            hc.alignment = Alignment(horizontal="center")

        # --- [합계행] A / B / C 총합 ---
        lbl = ws.cell(row=total_row, column=1, value="합계 (Total)")
        lbl.font = bold
        for cc, val in ((a_col, sum_a), (b_col, sum_b), (c_col, sum_c)):
            c = ws.cell(row=total_row, column=cc, value=val)
            c.number_format = money_fmt
            c.font = bold
        for cc in range(1, n_cols + 1):
            ws.cell(row=total_row, column=cc).border = thin_top

        # --- [청구액행] "총 청구액 = (A + B) - C" + 최종 금액(노란색) ---
        label_cell = ws.cell(row=bill_row, column=1, value="총 청구액 = (A + B) - C")
        label_cell.font = Font(bold=True, size=11)
        label_cell.alignment = Alignment(horizontal="left")
        # 라벨을 금액 셀 앞까지 병합해 한 줄로 보이게
        if c_col - 1 >= 2:
            try:
                ws.merge_cells(start_row=bill_row, start_column=1,
                               end_row=bill_row, end_column=c_col - 1)
            except Exception:
                pass

        amount_cell = ws.cell(row=bill_row, column=c_col, value=f"{final_amount:,}원")
        amount_cell.fill = PatternFill("solid", fgColor="FFFF00")   # 노란색 배경
        amount_cell.font = Font(bold=True, size=12)
        amount_cell.alignment = Alignment(horizontal="right")
        amount_cell.border = Border(top=Side(style="thin"), bottom=Side(style="thin"),
                                    left=Side(style="thin"), right=Side(style="thin"))

        # --- 컬럼 폭 자동(대략) ---
        for cc, name in enumerate(cols, start=1):
            sample = [str(name)]
            for rr in range(first_data, min(last_data, first_data + 200) + 1):
                v = ws.cell(row=rr, column=cc).value
                if v is not None:
                    sample.append(str(v))
            width = max(len(s) for s in sample) + 2
            ws.column_dimensions[get_column_letter(cc)].width = max(10, min(width, 42))

    return sum_a, sum_b, sum_c, final_amount


# =====================================================================
# 7. (선택) 테스트용 샘플 입력 파일 생성
# =====================================================================
def create_sample_input(path: str):
    """요구 서식과 동일한 2개 시트를 가진 샘플 엑셀을 만든다(테스트용)."""
    config = pd.DataFrame({
        COL_PRICE_ITEM:  ["에어컨 청소", "세탁기 분해청소", "냉장고 청소", "매트리스 케어"],
        COL_PRICE_VALUE: [80000, 120000, 100000, 90000],
    })
    # 컬럼 순서상 '작업분류(선택)' 이 8번째(H열)에 오도록 구성
    data = pd.DataFrame({
        "번호":      [1, 2, 3, 4, 5],
        "고객명":    ["김철수", "이영희", "박민수", "최지우", "정해인"],
        "필터1":     ["", "", "", "", ""],
        "필터2":     ["", "", "", "", ""],
        "필터3":     ["", "", "", "", "매트리스 케어"],   # 작업분류 빈칸 → 보조 스캔 매칭
        "연락처":    ["010-0000-0001", "010-0000-0002", "010-0000-0003",
                     "010-0000-0004", "010-0000-0005"],
        "방문일":    ["2026-05-02", "2026-05-05", "2026-05-11", "2026-05-18", "2026-05-24"],
        COL_CATEGORY: ["에어컨 청소", "세탁기 분해청소", "냉장고 청소", "알수없는항목", ""],
        COL_ADDRESS: [
            "인천광역시 미추홀구 인주대로 123",
            "인천광역시 중구 영종대로 100",        # '영종' → +11,000
            "서울특별시 강남구 테헤란로 500",       # 원거리 → 왕복 50km 이상 가능
            "인천광역시 연수구 컨벤시아대로 100",
            "인천광역시 중구 영종해안남로 321",     # '영종' → +11,000
        ],
        COL_FEE:    ["", "", "", "", ""],
        COL_TOLL:   ["", "", "", "", ""],
        COL_ROUND:  ["", "", "", "", ""],
        COL_ONSITE: [10000, 0, 5000, "", 20000],   # 빈칸 → 0 처리
    })
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        config.to_excel(w, sheet_name=SHEET_CONFIG, index=False)
        data.to_excel(w, sheet_name=SHEET_DATA, index=False)


# =====================================================================
# 8. 메인
# =====================================================================
def main():
    parser = argparse.ArgumentParser(description="렌탈/케어 정산 자동화")
    parser.add_argument("--make-sample", action="store_true",
                        help="테스트용 샘플 입력 엑셀을 생성하고 종료")
    parser.add_argument("--input",  default=INPUT_FILE,  help="입력 엑셀 경로")
    parser.add_argument("--output", default=OUTPUT_FILE, help="출력 엑셀 경로")
    args = parser.parse_args()

    input_file, output_file = args.input, args.output

    if args.make_sample:
        create_sample_input(input_file)
        print(f"✔ 샘플 입력 파일 생성 완료: {input_file}")
        return

    if not os.path.exists(input_file):
        print(f"[오류] 입력 파일을 찾을 수 없습니다: {input_file}")
        print("       --make-sample 로 샘플을 만들거나, --input 경로를 확인하세요.")
        sys.exit(1)

    print(f"[1/3] 단가표 로드 : {input_file} :: '{SHEET_CONFIG}'")
    price_map = load_price_map(input_file)
    print(f"      → 단가 항목 {len(price_map)}건")

    print(f"[2/3] 정산 계산  : '{SHEET_DATA}'")
    if not _api_ready():
        print("      [안내] KAKAO_API_KEY 미설정 → 수수료 + 영종 할증만 적용(거리 할증 생략)")
    df = process(input_file, price_map)

    print(f"[3/3] 결과 저장  : {output_file}")
    a, b, c, final_amount = save_with_summary(df, output_file)

    print("-" * 56)
    print(f"  수수료 총합      (A) : {a:>12,} 원")
    print(f"  추가통행료 총합  (B) : {b:>12,} 원")
    print(f"  현장결제 총합    (C) : {c:>12,} 원")
    print(f"  총 청구액  (A+B)-C   : {final_amount:>12,} 원")
    print("-" * 56)
    print("✔ 완료!")


if __name__ == "__main__":
    main()
